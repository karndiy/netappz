"""PDF to Word with good Thai + English results.

Two ways to get a .docx:
  * Text PDFs go through pdf2docx, which keeps the layout (columns, tables,
    pictures). Its output is then "polished" for Thai: Word draws Thai with the
    complex-script font, size, and bold settings (w:cs, w:szCs, w:bCs), which
    pdf2docx leaves out, so Thai would otherwise come out small and never bold.
  * Scanned PDFs (pictures of pages, or text that can't be decoded) are read
    with Tesseract OCR in Thai + English and rebuilt as editable paragraphs.

Thai text from older PDFs is also repaired: private-use glyph codes that some
Thai fonts use for shifted vowels and tone marks, a split sara am, and marks
stored in the wrong order.
"""

import re
import shutil
import statistics
import subprocess
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from converters import ConversionError, open_pdf

OCR_DPI = 300
MAX_OCR_PAGES = 80
THAI = re.compile("[\u0e00-\u0e7f]")
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# ------------------------------------------------------------------ Thai text

# Private-use glyphs (Adobe/Microsoft Thai fonts) -> the real Thai characters.
THAI_PUA = {
    0xF700: "\u0e10", 0xF701: "\u0e34", 0xF702: "\u0e35", 0xF703: "\u0e36", 0xF704: "\u0e37",
    0xF705: "\u0e48", 0xF706: "\u0e49", 0xF707: "\u0e4a", 0xF708: "\u0e4b", 0xF709: "\u0e4c",
    0xF70A: "\u0e48", 0xF70B: "\u0e49", 0xF70C: "\u0e4a", 0xF70D: "\u0e4b", 0xF70E: "\u0e4c",
    0xF70F: "\u0e0d", 0xF710: "\u0e31", 0xF711: "\u0e4d", 0xF712: "\u0e47", 0xF713: "\u0e48",
    0xF714: "\u0e49", 0xF715: "\u0e4a", 0xF716: "\u0e4b", 0xF717: "\u0e4c", 0xF718: "\u0e38",
    0xF719: "\u0e39", 0xF71A: "\u0e3a",
}
UPPER_VOWELS = "\u0e31\u0e34\u0e35\u0e36\u0e37\u0e47"
LOWER_VOWELS = "\u0e38\u0e39\u0e3a"
TONES = "\u0e48\u0e49\u0e4a\u0e4b\u0e4c"


def normalize_thai(text):
    """Repair Thai text extracted from PDFs so Word shows and searches it correctly."""
    if not text:
        return text
    text = text.translate(THAI_PUA)
    # Sara am stored as nikhahit + sara aa (with an optional tone mark in between).
    text = re.sub(f"\u0e4d([{TONES}]?)\u0e32", "\\1\u0e33", text)
    # A tone mark must follow the vowel above or below its consonant.
    text = re.sub(f"([{TONES}])([{UPPER_VOWELS}{LOWER_VOWELS}])", r"\2\1", text)
    # The same mark twice in a row comes from overlapping glyph runs.
    text = re.sub(f"([{UPPER_VOWELS}{LOWER_VOWELS}{TONES}])\\1+", r"\1", text)
    return text


# PDF font names -> the names Word users have installed.
FONT_NAMES = [
    (r"laksaman|thsarabunnew|th ?sarabun ?new", "TH Sarabun New"),
    (r"thsarabunpsk|th ?sarabun ?psk", "TH SarabunPSK"),
    (r"^sarabun", "Sarabun"),
    (r"thniramit", "TH Niramit AS"),
    (r"thchakrapetch", "TH Chakra Petch"),
    (r"thkodchasal", "TH Kodchasal"),
    (r"angsanaupc", "AngsanaUPC"),
    (r"angsana", "Angsana New"),
    (r"cordiaupc", "CordiaUPC"),
    (r"cordia", "Cordia New"),
    (r"browalliaupc", "BrowalliaUPC"),
    (r"browallia", "Browallia New"),
    (r"eucrosia", "EucrosiaUPC"),
    (r"leelawadeeui", "Leelawadee UI"),
    (r"leelawadee", "Leelawadee"),
    (r"notosansthailooped", "Noto Sans Thai Looped"),
    (r"notosansthai", "Noto Sans Thai"),
    (r"notoserifthai", "Noto Serif Thai"),
    (r"^(garuda|loma|umpush|waree|sawasdee)", "Tahoma"),
    (r"^(kinnari|norasi)", "Angsana New"),
    (r"^tlwg(typo|typist|mono)", "Courier New"),
    (r"arialmt|^arial$|^nimbussans|^liberationsans|^helvetica", "Arial"),
    (r"^nimbusroman|^liberationserif|^times$|^timesroman", "Times New Roman"),
    (r"^nimbusmono|^liberationmono|^courier$", "Courier New"),
    (r"^carlito", "Calibri"),
    (r"^caladea", "Cambria"),
    (r"timesnewromanps|^timesnewroman", "Times New Roman"),
    (r"couriernewps|^couriernew", "Courier New"),
]


def word_font(name):
    if not name:
        return name
    plain = re.sub(r"^[A-Z]{6}\+", "", name)  # Subset prefix, e.g. ABCDEF+Font.
    plain = re.sub(r"[-, ](bold|italic|oblique|regular|bolditalic|boldoblique|light|medium|mt|ps)+$", "", plain, flags=re.I)
    plain = re.sub(r"[-, ](bold|italic|oblique|regular)$", "", plain, flags=re.I)
    key = plain.replace(" ", "").replace("-", "").lower()
    for pattern, word in FONT_NAMES:
        if re.search(pattern, key) or re.search(pattern, plain.lower()):
            return word
    return plain


# ------------------------------------------------------------------ polishing

# Child order of w:rPr in the WordprocessingML schema; Word rejects files that break it.
RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow",
    "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
    "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs",
    "em", "lang", "eastAsianLayout", "specVanish", "oMath",
]


def _w(tag):
    return f"{{{W}}}{tag}"


def _set_child(parent, tag, attrs=None, order=RPR_ORDER):
    """Return parent's w:<tag>, creating it in schema order if needed."""
    existing = parent.find(_w(tag))
    if existing is None:
        existing = parent.makeelement(_w(tag), {})
        position = order.index(tag)
        later = {_w(name) for name in order[position + 1:]}
        for index, child in enumerate(parent):
            if child.tag in later:
                parent.insert(index, existing)
                break
        else:
            parent.append(existing)
    for key, value in (attrs or {}).items():
        existing.set(_w(key), value)
    return existing


def _on(element):
    return element is not None and element.get(_w("val"), "1") not in ("0", "false", "off")


def _polish_part(xml):
    from lxml import etree

    root = etree.fromstring(xml)
    thai_fonts = {}
    for run in root.iter(_w("r")):
        texts = run.findall(_w("t"))
        if not texts:
            continue
        for node in texts:
            node.text = normalize_thai(node.text)
        has_thai = any(THAI.search(node.text or "") for node in texts)
        props = run.find(_w("rPr"))
        if props is None:
            if not has_thai:
                continue
            props = run.makeelement(_w("rPr"), {})
            run.insert(0, props)

        fonts = props.find(_w("rFonts"))
        if fonts is not None:
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                if fonts.get(_w(key)):
                    fonts.set(_w(key), word_font(fonts.get(_w(key))))
        if not has_thai:
            continue

        # Thai is a complex script in Word: give it the same font, size and style.
        font = (fonts.get(_w("cs")) or fonts.get(_w("ascii")) or fonts.get(_w("hAnsi"))) if fonts is not None else None
        font = font or "TH Sarabun New"
        fonts = _set_child(props, "rFonts", {"cs": font})
        for key in ("ascii", "hAnsi"):
            if not fonts.get(_w(key)):
                fonts.set(_w(key), font)
        if _on(props.find(_w("b"))):
            _set_child(props, "bCs")
        if _on(props.find(_w("i"))):
            _set_child(props, "iCs")
        size = props.find(_w("sz"))
        if size is not None:
            _set_child(props, "szCs", {"val": size.get(_w("val"))})
            thai_fonts[font] = thai_fonts.get(font, 0) + len("".join(node.text or "" for node in texts))
        _set_child(props, "lang", {"bidi": "th-TH"})

    # Exact line spacing clips Thai tone marks when Word substitutes a taller font.
    for paragraph in root.iter(_w("p")):
        if not THAI.search("".join(paragraph.itertext())):
            continue
        spacing = paragraph.find(f"{_w('pPr')}/{_w('spacing')}")
        if spacing is not None and spacing.get(_w("lineRule")) == "exact":
            spacing.set(_w("lineRule"), "atLeast")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True), thai_fonts


def _polish_styles(xml, thai_font):
    from lxml import etree

    root = etree.fromstring(xml)
    defaults = root.find(_w("docDefaults"))
    if defaults is None:
        defaults = root.makeelement(_w("docDefaults"), {})
        root.insert(0, defaults)
    rpr_default = defaults.find(_w("rPrDefault"))
    if rpr_default is None:
        rpr_default = defaults.makeelement(_w("rPrDefault"), {})
        defaults.insert(0, rpr_default)
    props = rpr_default.find(_w("rPr"))
    if props is None:
        props = rpr_default.makeelement(_w("rPr"), {})
        rpr_default.append(props)
    # New Thai text typed into the document uses the document's Thai font.
    fonts = _set_child(props, "rFonts", {"cs": thai_font})
    if fonts.get(_w("csTheme")) is not None:
        del fonts.attrib[_w("csTheme")]
    _set_child(props, "lang", {"bidi": "th-TH"})
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def polish_docx(path):
    """Fix Thai fonts, sizes, bold, and text inside a .docx, in place."""
    path = Path(path)
    temp = path.with_suffix(".polished.docx")
    thai_fonts = {}
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        parts = {}
        for item in source.infolist():
            data = source.read(item.filename)
            if re.fullmatch(r"word/(document|header\d*|footer\d*|footnotes|endnotes)\.xml", item.filename):
                data, fonts = _polish_part(data)
                for font, count in fonts.items():
                    thai_fonts[font] = thai_fonts.get(font, 0) + count
            parts[item.filename] = (item, data)
        if "word/styles.xml" in parts and thai_fonts:
            item, data = parts["word/styles.xml"]
            parts["word/styles.xml"] = (item, _polish_styles(data, max(thai_fonts, key=thai_fonts.get)))
        for item, data in parts.values():
            target.writestr(item, data)
    temp.replace(path)
    return path


# ------------------------------------------------------------------ choosing a method

def _letters(text):
    return len(re.findall(r"[A-Za-z\u0e01-\u0e2e\u0e30-\u0e39\u0e40-\u0e4e0-9]", text))


def page_needs_ocr(page):
    """True when a page is a picture of text, or its text can't be decoded."""
    raw = page.get_text()
    text = normalize_thai(raw)
    letters = _letters(text)
    if letters >= 30:
        # Text is there, but is it readable? Broken font encodings give control
        # characters, replacement marks, or unknown private-use codes.
        bad = len(re.findall("[\ufffd\x00-\x08\x0e-\x1f\ue000-\uf8ff]", text))
        return bad > letters * 0.25
    area = abs(page.rect)
    covered = 0
    for image in page.get_image_info():
        covered += abs(page.rect & image["bbox"])
    return covered > area * 0.4


def ocr_available():
    tesseract = shutil.which("tesseract")
    if not tesseract:
        return False
    try:
        langs = subprocess.run([tesseract, "--list-langs"], capture_output=True, text=True, timeout=20).stdout
    except (OSError, subprocess.SubprocessError):
        return False
    return "tha" in langs.split() and "eng" in langs.split()


# ------------------------------------------------------------------ OCR

# Tesseract's "best" models read Thai far better than the small ones Debian ships;
# the Docker image installs them here (see Dockerfile).
BEST_TESSDATA = Path("/usr/share/tesseract-ocr/best")


def _ocr_words(image_path, langs, psm=3, best=False):
    """Run Tesseract and return its words in reading order.

    image_path may be a .txt file listing several images: Tesseract then loads its
    model once for all of them, and each word's "image" says which one it came from.
    """
    command = ["tesseract", str(image_path), "stdout", "-l", langs, "--oem", "1", "--psm", str(psm),
               "--dpi", str(OCR_DPI), "-c", "preserve_interword_spaces=1"]
    if best and (BEST_TESSDATA / "tha.traineddata").exists():
        command += ["--tessdata-dir", str(BEST_TESSDATA)]
    result = subprocess.run(
        command + ["tsv"],
        capture_output=True, text=True, timeout=600,
        # One thread per Tesseract: several threads fight over the CPU and run ~8x slower.
        env={"OMP_THREAD_LIMIT": "1", "PATH": "/usr/bin:/bin:/usr/local/bin"},
    )
    if result.returncode != 0:
        raise ConversionError("Text recognition (OCR) failed on a page.")
    words = []
    for row in result.stdout.splitlines()[1:]:
        parts = row.split("\t")
        if len(parts) < 12 or parts[0] != "5" or not parts[11].strip():
            continue
        try:
            conf = float(parts[10])
        except ValueError:
            conf = 0.0
        words.append({
            "image": int(parts[1]) - 1,
            "line": (int(parts[2]), int(parts[3]), int(parts[4])),
            "left": int(parts[6]), "top": int(parts[7]), "width": int(parts[8]), "height": int(parts[9]),
            "conf": conf, "text": parts[11].strip(),
        })
    return words


def _clean_scan(image_path):
    """Remove what OCR misreads on forms, and remember it to put back as text.

    Dotted fill-in lines turn into strings of digits, table rules into letters, and
    tick boxes into stray characters. They are erased from a cleaned copy of the page
    before OCR; their positions come back as fill tokens ("....") and ☐ / ○ marks.
    Returns (path to read, extra tokens).
    """
    try:
        import cv2
        import numpy as np
    except ImportError:
        return image_path, []
    gray = cv2.imread(str(image_path), cv2.IMREAD_GRAYSCALE)
    _, ink = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    clean = ink.copy()
    # Long straight rules: table borders, boxes, underlines.
    for size in ((80, 1), (1, 80)):
        rules = cv2.morphologyEx(ink, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, size))
        clean[cv2.dilate(rules, np.ones((3, 3), np.uint8)) > 0] = 0

    tokens = []
    count, labels, stats, centers = cv2.connectedComponentsWithStats(clean, 8)
    # Dotted lines: 5 or more dot-sized blobs on one baseline, close together.
    dots = sorted(
        (i for i in range(1, count) if stats[i, 2] <= 10 and stats[i, 3] <= 10 and stats[i, 4] <= 60),
        key=lambda i: (int(centers[i][1] // 6), centers[i][0]),
    )
    rows = {}
    for i in dots:
        rows.setdefault(int(centers[i][1] // 6), []).append(i)
    used = set()
    erase = np.zeros(count, bool)
    for key in sorted(rows):
        run = []
        for i in sorted((i for i in rows[key] + rows.get(key + 1, []) if i not in used), key=lambda i: centers[i][0]):
            if run and (centers[i][0] - centers[run[-1]][0] > 30 or abs(centers[i][1] - centers[run[-1]][1]) > 5):
                run = run if len(run) >= 5 else []
                if run:
                    used.update(run)
                    tokens.append(("fill", run))
                run = []
            run.append(i)
        if len(run) >= 5:
            used.update(run)
            tokens.append(("fill", run))
    extras = []
    for _, run in tokens:
        erase[run] = True
        left = min(stats[i, 0] for i in run)
        right = max(stats[i, 0] + stats[i, 2] for i in run)
        top = min(stats[i, 1] for i in run)
        bottom = max(stats[i, 1] + stats[i, 3] for i in run)
        extras.append({"kind": "fill", "left": left, "top": top, "width": right - left, "height": bottom - top})
    # Tick boxes and circles: hollow, about a letter tall, as wide as tall.
    for i in range(1, count):
        x, y, w, h, area = stats[i]
        if erase[i] or not (24 <= w <= 50 and 24 <= h <= 50 and 0.85 <= w / h <= 1.18):
            continue
        middle = clean[y + h // 4: y + 3 * h // 4, x + w // 4: x + 3 * w // 4]
        if area / (w * h) < 0.45 and middle.mean() < 20:
            square = clean[y:y + h // 5, x:x + w // 5].mean() > 60  # Ink in the corner: a box, not a circle.
            erase[i] = True
            extras.append({"kind": "box", "text": "\u2610" if square else "\u25cb",
                           "left": x, "top": y, "width": w, "height": h})
    clean[erase[labels]] = 0
    path = image_path.with_name(image_path.stem + "-clean.png")
    cv2.imwrite(str(path), 255 - clean)
    return path, extras


LATIN_WORD = re.compile(r"[\x21-\x7e\u2018\u2019\u201c\u201d\u2013\u2014]*[A-Za-z0-9][\x21-\x7e\u2018\u2019\u201c\u201d\u2013\u2014]*")


def _merge_english(line, english):
    """Thai+English OCR misreads English inside Thai lines; use the English-only pass there.

    The English pass reads Thai as low-confidence nonsense, so only its confident,
    Latin-looking words replace the Thai pass's words under them.
    """
    top = min(word["top"] for word in line)
    bottom = max(word["top"] + word["height"] for word in line)
    candidates = []
    for word in english:
        middle = word["top"] + word["height"] / 2
        if word["conf"] >= 85 and top <= middle <= bottom and LATIN_WORD.fullmatch(word["text"]):
            candidates.append(word)
    def covered(english_word):
        return [w for w in line if english_word["left"] <= w["left"] + w["width"] / 2 <= english_word["left"] + english_word["width"]]

    def trusted(english_word):
        # Numbers and real words win. One or two letters over confident Thai are usually
        # the English pass turning Thai letters into look-alikes ("a", "ly", "WA").
        if re.search(r"\d", english_word["text"]) or len(re.findall(r"[A-Za-z]", english_word["text"])) >= 3:
            return True
        under = covered(english_word)
        return not any(THAI.search(w["text"]) for w in under) or statistics.mean(w["conf"] for w in under) < 60

    candidates = [e for e in candidates if trusted(e)]
    tokens = []
    used = set()
    for word in line:
        center = word["left"] + word["width"] / 2
        cover = next((i for i, e in enumerate(candidates) if e["left"] <= center <= e["left"] + e["width"]), None)
        if cover is None:
            tokens.append(word)
        elif cover not in used:
            used.add(cover)
            tokens.append(candidates[cover])
    return tokens


MARKS = re.compile("[\u0e31\u0e34-\u0e3a\u0e47-\u0e4e]")
SPACE_GAP = 0.5  # Gap between Thai glyphs, in consonant heights, that counts as a space.


def _unit(words):
    """Height of a plain Thai consonant (or Latin letters) on this line, in pixels."""
    tallest = max(w["height"] for w in words)
    # Some consonant boxes include the marks above and below; leave those out.
    consonants = [w["height"] for w in words if re.fullmatch("[\u0e01-\u0e2e]", w["text"]) and w["height"] < tallest * 0.75]
    if consonants:
        return statistics.median(consonants)
    ems = [w["height"] / _latin_ratio(w["text"]) for w in words if LATIN_WORD.fullmatch(w["text"])]
    if ems:
        return statistics.median(ems) * 0.61
    return statistics.median(w["height"] for w in words) * 0.5


ASCENDERS = set("ABCDEFGHIJKLMNOPQRSTUVWXYZbdfhklt0123456789%()[]{}/\\|!?\"'&#$@")
DESCENDERS = set("gjpqyQ,;()[]{}@")


def _latin_ratio(text):
    """How tall a Latin word's box is, as a share of the font size."""
    up = any(c in ASCENDERS for c in text)
    down = any(c in DESCENDERS for c in text)
    return 0.95 if up and down else 0.76 if up else 0.74 if down else 0.54


def _line_text(tokens, unit):
    """Join tokens using the real gaps: Thai has no spaces between words, so gaps decide."""
    text = ""
    right = None
    for token in tokens:
        if token.get("kind") == "fill":
            # Dots about as long as the original line: a "." is roughly 0.27 em wide.
            piece = "." * max(3, round(token["width"] / (unit / 0.61 * 0.27)))
        else:
            piece = normalize_thai(token["text"])
        if right is not None and not MARKS.match(piece):  # Never a space before a vowel or tone mark.
            gap = token["left"] - right
            thai_before, thai_after = THAI.search(text[-1:]), THAI.search(piece[:1])
            # Tesseract splits English at spaces already; Thai comes one glyph at a time.
            limit = SPACE_GAP if thai_before and thai_after else 0.25 if thai_before or thai_after else 0.1
            if gap > unit * 4:
                text += "\t"  # Columns, such as a table row.
            elif gap > unit * limit:
                text += " "
        text += piece
        end = token["left"] + token["width"]
        right = end if right is None else max(right, end)
    # Repair across tokens too: OCR returns a sara am's two parts as separate glyphs.
    return normalize_thai(text)


def _font_size(unit):
    # A plain Thai consonant is about 0.61 em tall.
    return unit * 72 / OCR_DPI / 0.61


def _ink(image, words):
    """Share of dark pixels inside the glyph boxes; bold text has more."""
    dark = total = 0
    for word in words:
        box = image.crop((word["left"], word["top"], word["left"] + word["width"], word["top"] + word["height"]))
        histogram = box.histogram()
        dark += sum(histogram[:128])
        total += sum(histogram)
    return dark / total if total else 0


def _box(words):
    return (min(w["left"] for w in words), min(w["top"] for w in words),
            max(w["left"] + w["width"] for w in words), max(w["top"] + w["height"] for w in words))


def _reread_lines(image_path, words):
    """Fix Thai lines that lost their tone marks; returns [(top, bottom, words)] in reading order.

    On a page, Tesseract often puts the row of vowels and tone marks above a Thai line
    into a "line" of its own: that row reads as junk and the real line loses its marks
    (เรียน comes out as เรยน). Those rows are merged back into the line they belong to,
    and each line is read again on its own (one text line per image), which keeps the
    marks with their letters.
    """
    from PIL import Image

    grouped = {}
    for word in words:
        grouped.setdefault(word["line"], []).append(word)
    lines = []
    for key, items in grouped.items():  # Tesseract's reading order.
        left, top, right, bottom = _box(items)
        lines.append({"key": key, "words": items, "left": left, "top": top, "right": right, "bottom": bottom,
                      "height": statistics.median(w["height"] for w in items),
                      "chars": sum(len(w["text"]) for w in items),
                      "conf": statistics.mean(w["conf"] for w in items)})
    if not lines:
        return []
    # Tesseract returns whole words for clean prints but single glyphs for many scans, so count characters.
    typical = statistics.median(l["height"] for l in lines if l["chars"] >= 5) if any(
        l["chars"] >= 5 for l in lines) else statistics.median(l["height"] for l in lines)
    spacing = statistics.median(l["bottom"] - l["top"] for l in lines)

    def real(line):
        # Rows of marks have short glyph boxes and read as unsure junk; a real line with
        # fill-in dots can be unsure too, but its letters are full height.
        return line["height"] >= typical * 0.8 and line["conf"] >= 55 and line["chars"] >= 3

    kept = [line for line in lines if real(line)]
    for line in lines:
        if real(line):
            continue
        near = None
        for other in kept:
            if not real(other) or min(other["right"], line["right"]) <= max(other["left"], line["left"]):
                continue
            # Overlapping it, just above it (upper marks), or just below it (lower marks).
            gap = max(other["top"] - line["bottom"], line["top"] - other["bottom"])
            if gap < spacing * 0.5 and (near is None or gap < near[0]):
                near = (gap, other)
        if near:
            other = near[1]
            other["top"], other["bottom"] = min(other["top"], line["top"]), max(other["bottom"], line["bottom"])
        else:
            kept.append(line)  # A short line on its own, such as a label.
    kept.sort(key=lambda line: line["key"])

    # Read every line again, one line per image, in one Tesseract run per CPU.
    page = Image.open(image_path)
    ink = page.convert("L").point(lambda v: 255 if v < 128 else 0)
    pad = 10
    reach = int(typical * 2.5)  # Letters further than this from the line belong to something else.
    for line in kept:
        # Tesseract sometimes misses the first or last letters; take in ink that continues the line.
        band = ink.crop((0, line["top"], page.width, line["bottom"]))
        columns = [x for x, value in enumerate(band.resize((page.width, 1), Image.BOX).getdata()) if value]
        for x in reversed([c for c in columns if c < line["left"]]):
            if line["left"] - x > reach:
                break
            line["left"] = x
        for x in [c for c in columns if c > line["right"]]:
            if x - line["right"] > reach:
                break
            line["right"] = x
    _add_missed_ink(image_path, kept, typical, reach)
    # Lines that now overlap are one line read in pieces; read it once.
    merged = []
    for line in kept:
        same = next((m for m in merged if min(m["right"], line["right"]) > max(m["left"], line["left"])
                     and min(m["bottom"], line["bottom"]) - max(m["top"], line["top"])
                     > min(m["bottom"] - m["top"], line["bottom"] - line["top"]) * 0.5), None)
        if same:
            for side, pick in (("left", min), ("top", min), ("right", max), ("bottom", max)):
                same[side] = pick(same[side], line[side])
            same["words"] = same["words"] + line["words"]
            same["conf"] = statistics.mean(w["conf"] for w in same["words"])
        else:
            merged.append(line)
    kept = merged
    crops = []
    for number, line in enumerate(kept):
        origin = (max(0, line["left"] - pad), max(0, line["top"] - pad))
        crop = page.crop((*origin, min(page.width, line["right"] + pad), min(page.height, line["bottom"] + pad)))
        path = image_path.with_name(f"{image_path.stem}-line{number:03d}.png")
        crop.save(path)
        crops.append((path, origin))
    halves = [list(range(len(crops)))[i::2] for i in range(2)]

    def read(indexes):
        if not indexes:
            return []
        listing = image_path.with_name(f"{image_path.stem}-lines{indexes[0] % 2}.txt")
        listing.write_text("".join(f"{crops[i][0]}\n" for i in indexes))
        found = _ocr_words(listing, "tha+eng", psm=7, best=True)
        for word in found:
            word["image"] = indexes[word["image"]]
        return found

    with ThreadPoolExecutor(max_workers=2) as pool:
        again = [word for part in pool.map(read, halves) for word in part]
    by_line = {}
    for word in again:
        x, y = crops[word["image"]][1]
        word["left"] += x
        word["top"] += y
        by_line.setdefault(word["image"], []).append(word)

    result = []
    for number, line in enumerate(kept):
        new = by_line.get(number, [])
        # Keep the first reading when the second is clearly less sure, which happens when
        # a line crop catches pieces of its neighbours.
        if new and statistics.mean(w["conf"] for w in new) >= line["conf"] - 5:
            chosen = new
        else:
            chosen = line["words"]
        if not chosen:
            continue
        # Leftover rows of marks and specks (stamps, emblems) read as unsure junk.
        if statistics.mean(w["conf"] for w in chosen) < 70:
            continue
        result.append((line["top"], line["bottom"], chosen))
    return result


def _add_missed_ink(image_path, lines, typical, reach):
    """Tesseract's page layout sometimes skips a word standing alone, such as "สัญชาติ"
    between two erased fill-in lines. Find letter-sized ink outside every line and add it:
    to the line on the same row, or as a line of its own."""
    try:
        import cv2
    except ImportError:
        return
    gray = cv2.imread(str(image_path), cv2.IMREAD_GRAYSCALE)
    _, ink = cv2.threshold(gray, 128, 255, cv2.THRESH_BINARY_INV)
    count, _, stats, _ = cv2.connectedComponentsWithStats(ink, 8)
    pad = typical * 0.3
    loose = []
    for i in range(1, count):
        x, y, w, h, _ = stats[i]
        if h < typical * 0.5 or h > typical * 3 or w > typical * 6:
            continue  # Specks, marks, pictures, rules.
        cx, cy = x + w / 2, y + h / 2
        if any(l["left"] - pad <= cx <= l["right"] + pad and l["top"] - pad <= cy <= l["bottom"] + pad for l in lines):
            continue
        loose.append((x, y, x + w, y + h))
    # Group into runs of letters: same row, close together.
    groups = []
    for box in sorted(loose):
        group = next((g for g in groups if box[0] - g[2] < reach
                      and min(g[3], box[3]) - max(g[1], box[1]) > min(g[3] - g[1], box[3] - box[1]) * 0.5), None)
        if group:
            group[:] = [min(group[0], box[0]), min(group[1], box[1]), max(group[2], box[2]), max(group[3], box[3])]
        else:
            groups.append(list(box))
    for left, top, right, bottom in groups:
        if right - left < typical * 0.8:
            continue  # A single stray blob.
        row = next((l for l in lines if min(l["bottom"], bottom) - max(l["top"], top) > (bottom - top) * 0.5), None)
        if row:
            row["left"], row["right"] = min(row["left"], left), max(row["right"], right)
            row["top"], row["bottom"] = min(row["top"], top), max(row["bottom"], bottom)
        else:
            above = [l for l in lines if l["top"] <= top]
            key = (max(above, key=lambda l: l["top"])["key"] if above else (0,)) + (1,)
            lines.append({"key": key, "words": [], "left": left, "top": top, "right": right, "bottom": bottom,
                          "conf": 0.0, "height": bottom - top, "chars": 0})
    lines.sort(key=lambda line: line["key"])


def _place_extras(lines, extras):
    """Add fill lines and tick boxes to the text line they sit on (or drop them)."""
    for extra in extras:
        middle = extra["top"] + extra["height"] / 2
        best = None
        for index, (top, bottom, words) in enumerate(lines):
            height = bottom - top
            if not (top - height * 0.2 <= middle <= bottom + height * 0.3):
                continue
            left, _, right, _ = _box(words)
            distance = max(left - (extra["left"] + extra["width"]), extra["left"] - right, 0)
            if best is None or distance < best[0]:
                best = (distance, index)
        if best:
            lines[best[1]][2].append({**extra, "conf": 100, "text": extra.get("text", "")})


def _ocr_page(image_path):
    from PIL import Image

    clean_path, extras = _clean_scan(image_path)
    found = _ocr_words(clean_path, "tha+eng")
    english = _ocr_words(clean_path, "eng") if found else []
    found = [w for w in found if not re.fullmatch(r"[|\[\]]+", w["text"])]  # Table borders read as characters.
    text_lines = _reread_lines(clean_path, found)
    _place_extras(text_lines, extras)
    image = Image.open(image_path).convert("L")

    lines = []
    for _, _, words in text_lines:
        letters = [w for w in words if "kind" not in w]
        if not letters:
            continue
        unit = _unit(letters)
        # Keep Tesseract's order for the letters: Thai vowel and tone mark boxes overlap
        # their neighbours, so sorting by position would scramble them.
        tokens = _merge_english(letters, english)
        for extra in sorted((w for w in words if "kind" in w), key=lambda w: w["left"]):
            at = next((i for i, t in enumerate(tokens) if t["left"] >= extra["left"]), len(tokens))
            tokens.insert(at, extra)
        first = tokens[0]
        left, top, right, bottom = _box(words)
        lines.append({
            "text": _line_text(tokens, unit),
            # Width of the first word: Thai comes a glyph at a time, so assume a short word.
            "first_width": first["width"] if LATIN_WORD.fullmatch(first.get("text", "")) else unit * 3,
            "left": left, "right": right, "top": top, "bottom": bottom,
            "size": _font_size(unit),
            "ink": _ink(image, letters),
            "conf": statistics.mean(w["conf"] for w in letters),
            "form": any("kind" in w for w in words),
        })
    if not lines:
        return []

    # Rebuild paragraphs: Tesseract often makes every line its own paragraph.
    column_left = min(line["left"] for line in lines)
    column_right = max(line["right"] for line in lines)
    column = max(column_right - column_left, 1)
    width = image.width
    for line in lines:
        middle = (line["left"] + line["right"]) / 2
        line["center"] = (
            line["right"] - line["left"] < column * 0.8
            and abs(middle - (column_left + column_right) / 2) < column * 0.04
            and line["left"] - column_left > column * 0.08
        )
    paragraphs = []
    for line in lines:
        previous = paragraphs[-1]["lines"][-1] if paragraphs else None
        joins = False
        # Form lines (fill-in dots, tick boxes) stay on their own, like the printed form.
        if previous and "\t" not in line["text"] and "\t" not in previous["text"] and not (line["form"] or previous["form"]):
            pitch = line["top"] - previous["top"]
            same_size = abs(line["size"] - previous["size"]) < previous["size"] * 0.2
            # The previous line was full if this line's first word wouldn't have fit after it.
            full = column_right - previous["right"] < line["first_width"] + column * 0.04
            same_start = abs(line["left"] - column_left) < column * 0.03
            joins = (
                # Lines of a paragraph sit about 2 font sizes apart; paragraphs further.
                0 < pitch < max(line["size"], previous["size"]) * OCR_DPI / 72 * 2.35
                and same_size
                and not previous["center"] and not line["center"]
                and full and same_start
            )
        if joins:
            paragraphs[-1]["lines"].append(line)
        else:
            paragraphs.append({"lines": [line]})

    result = []
    for paragraph in paragraphs:
        group = paragraph["lines"]
        result.append({
            "text": _join_lines([line["text"] for line in group]),
            "size": statistics.median(line["size"] for line in group),
            "ink": statistics.mean(line["ink"] for line in group),
            "center": all(line["center"] for line in group),
            "left": group[0]["left"] if len(group) == 1 else min(line["left"] for line in group[1:]),
            "first": group[0]["left"],
            "top": group[0]["top"],
            "confidence": statistics.mean(line["conf"] for line in group),
            "page_width": width,
        })
    return result


def _join_lines(lines):
    text = lines[0]
    for line in lines[1:]:
        thai_join = THAI.search(text[-1:]) and THAI.search(line[:1])
        if text.endswith("-") and not thai_join:
            text = text[:-1] + line
        else:
            text += ("" if thai_join else " ") + line
    return text


def ocr_to_docx(source, workdir, result):
    import pymupdf
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Emu, Pt

    document = open_pdf(source)
    try:
        if document.page_count > MAX_OCR_PAGES:
            raise ConversionError(f"Text recognition handles up to {MAX_OCR_PAGES} pages at a time.")
        images = []
        sizes = []
        for number, page in enumerate(document):
            image = workdir / f"ocr-{number:03d}.png"
            page.get_pixmap(dpi=OCR_DPI, colorspace=pymupdf.csGRAY).save(image)
            images.append(image)
            sizes.append((page.rect.width, page.rect.height))
    finally:
        document.close()

    with ThreadPoolExecutor(max_workers=2) as pool:
        pages = list(pool.map(_ocr_page, images))

    doc = Document()
    section = doc.sections[0]
    width, height = sizes[0]
    section.page_width, section.page_height = Emu(int(width * 12700)), Emu(int(height * 12700))
    margin = Emu(int(54 * 12700))
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = margin
    normal = doc.styles["Normal"]
    normal.font.name = "TH Sarabun New"
    normal.font.size = Pt(16)

    every = [p for page in pages for p in page]
    body_size = statistics.median(p["size"] for p in every) if every else 16
    body_ink = statistics.median(p["ink"] for p in every) if every else 0
    recognized = 0
    for number, (paragraphs, image) in enumerate(zip(pages, images)):
        if number:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if sum(len(p["text"]) for p in paragraphs) < 5:
            # Nothing readable, such as a photo: keep the page as a picture.
            doc.add_picture(str(image), width=section.page_width - section.left_margin - section.right_margin)
            continue
        for item in paragraphs:
            recognized += 1
            size = item["size"]
            # Body text gets one consistent size; headings keep theirs, rounded to half points.
            size = body_size if abs(size - body_size) < body_size * 0.2 else size
            size = round(min(max(size, 8), 48) * 2) / 2
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(item["text"])
            run.font.size = Pt(size)
            run.bold = size >= body_size * 1.25 or (body_ink and item["ink"] > body_ink * 1.25)
            if item["center"]:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                indent = item["left"] / OCR_DPI * 72 - 54
                if indent > 18:
                    paragraph.paragraph_format.left_indent = Pt(min(indent, width / 2))
            paragraph.paragraph_format.space_after = Pt(4)
    if not recognized:
        raise ConversionError("No text could be recognized in this PDF.")
    doc.save(result)
    return result


# ------------------------------------------------------------------ entry point

def pdf_to_word(source, workdir, ocr="auto"):
    """Convert a PDF to .docx. Returns (path, note) where note describes what was done."""
    from pdf2docx import Converter

    document = open_pdf(source)
    try:
        scanned = sum(1 for page in document if page_needs_ocr(page))
        total = document.page_count
    finally:
        document.close()

    use_ocr = ocr == "on" or (ocr == "auto" and scanned and scanned >= total / 2)
    result = workdir / f"{source.stem}.docx"
    if use_ocr:
        if not ocr_available():
            if ocr == "on":
                raise ConversionError(
                    "Text recognition (OCR) needs Tesseract with Thai, which isn't installed on this server. "
                    "Run the app with Docker to use it."
                )
            use_ocr = False
        else:
            ocr_to_docx(source, workdir, result)
            polish_docx(result)
            return result, f"Text recognized with OCR (Thai + English) on {total} page{'s' * (total != 1)}. Check names and numbers."

    try:
        converter = Converter(str(source))
        try:
            converter.convert(str(result))
        finally:
            converter.close()
    except Exception:
        raise ConversionError("Could not convert this PDF to Word.")
    polish_docx(result)
    note = ""
    if scanned and ocr == "off":
        note = f"{scanned} page{'s' * (scanned != 1)} look scanned and stay as pictures."
    elif scanned:
        note = f"{scanned} scanned page{'s' * (scanned != 1)} kept as pictures; choose “Always” to recognize their text."
    return result, note
